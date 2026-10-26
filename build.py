# -*- coding: utf-8 -*-
"""从《技能展厅.xlsx》生成静态站点到 dist/。

用法:
    python build.py

维护方式:
    只改 技能展厅.xlsx，不要手改 dist/ 里的文件（每次生成都会被覆盖）。
"""

from __future__ import annotations

import json
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    print("[ERROR] 缺少 openpyxl，请先安装:  pip install openpyxl")
    sys.exit(1)

ROOT = Path(__file__).resolve().parent
XLSX = ROOT / "技能展厅.xlsx"
SRC = ROOT / "src"
DIST = ROOT / "dist"

SHEET_CONFIG = "站点配置"
SHEET_SKILLS = "技能清单"
SHEET_CASES = "价值案例"
SHEET_GUIDE = "使用指南"

errors: list[str] = []
warnings: list[str] = []


# ---------------------------------------------------------------- helpers
def cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def to_int(value, field: str, where: str) -> int:
    raw = cell_str(value)
    if raw == "":
        warnings.append(f"{where}: 《{field}》为空，按 0 处理")
        return 0
    try:
        return int(round(float(raw)))
    except ValueError:
        errors.append(f"{where}: 《{field}》不是数字 -> {raw!r}")
        return 0


def to_float(value, field: str, where: str) -> float:
    raw = cell_str(value)
    if raw == "":
        return 0.0
    try:
        return float(raw)
    except ValueError:
        errors.append(f"{where}: 《{field}》不是数字 -> {raw!r}")
        return 0.0


def gain_label(minutes: int) -> str:
    """把分钟数变成「单次省 X 小时 / X 分钟」。"""
    if minutes <= 0:
        return ""
    if minutes < 60:
        return f"单次省 {minutes} 分钟"
    hours = minutes / 60
    text = f"{hours:.1f}".rstrip("0").rstrip(".")
    return f"单次省 {text} 小时"


def split_multi(text: str) -> list[str]:
    """把一格里用单独一行 --- 分隔的多段内容拆开。"""
    parts = re.split(r"(?m)^\s*[-—–]{3,}\s*$", text)
    return [p.strip() for p in parts if p.strip()]


def split_tags(text: str) -> list[str]:
    return [t.strip() for t in re.split(r"[,，、;；/\s]+", text) if t.strip()]


def read_rows(ws, label: str) -> list[dict]:
    headers = [cell_str(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    if not any(headers):
        errors.append(f"{label}: 第一行表头是空的")
        return []
    rows = []
    for idx, raw in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(v is None or str(v).strip() == "" for v in raw):
            continue
        row = {}
        for i, name in enumerate(headers):
            if name and i < len(raw):
                row[name] = raw[i]
        rows.append({"__row__": idx, **row})
    if not rows:
        errors.append(f"{label}: 没有读到任何数据行")
    return rows


def need(row: dict, field: str, label: str) -> str:
    value = cell_str(row.get(field))
    if not value:
        errors.append(f"{label} 第{row['__row__']}行: 缺少《{field}》")
    return value


def warn_missing_columns(ws, label: str, expected: list[str]) -> None:
    headers = [cell_str(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    unknown = [h for h in headers if h and h not in expected]
    missing = [h for h in expected if h not in headers]
    if missing:
        warnings.append(f"{label}: 缺少列 {missing}（将按空值处理）")
    if unknown:
        warnings.append(f"{label}: 忽略未知列 {unknown}")


# ---------------------------------------------------------------- load
if not XLSX.exists():
    print(f"[ERROR] 找不到数据文件: {XLSX}")
    sys.exit(1)

wb = openpyxl.load_workbook(XLSX, data_only=True)
for name in (SHEET_CONFIG, SHEET_SKILLS, SHEET_CASES, SHEET_GUIDE):
    if name not in wb.sheetnames:
        errors.append(f"工作簿缺少 Sheet: {name}")
if errors:
    print("[ERROR] " + "\n[ERROR] ".join(errors))
    sys.exit(1)

# --- 站点配置
cfg: dict[str, str] = {}
for row in read_rows(wb[SHEET_CONFIG], SHEET_CONFIG):
    key = cell_str(row.get("配置项"))
    if key:
        cfg[key] = cell_str(row.get("值"))

defaults = {
    "站点名称": "Skill 展厅",
    "主标题": "把专业活儿做完",
    "描述": "",
    "角标": "智能体平台 · Skill 资产库",
    "团队署名": "",
    "主题色": "#2563eb",
    "版本": "v1.0",
    "联系方式": "",
    "页脚": "由 Excel 一键生成",
    "关于": "",
}
for k, v in defaults.items():
    if not cfg.get(k):
        cfg[k] = v
if not re.fullmatch(r"#[0-9a-fA-F]{6}", cfg["主题色"]):
    warnings.append(f"主题色 {cfg['主题色']!r} 不是 #RRGGBB，已回退为 #2563eb")
    cfg["主题色"] = "#2563eb"

# --- 技能清单
SKILL_COLS = ["编号", "技能名称", "英文名称", "一句话简介", "业务分类", "功能分类", "业务负责人",
              "IT负责人", "适用场景", "提示词示例", "单次省时(分钟)", "每周次数",
              "提升价值", "标签", "状态", "推荐位", "排序权重"]
warn_missing_columns(wb[SHEET_SKILLS], SHEET_SKILLS, SKILL_COLS)

skills = []
seen_ids = set()
for row in read_rows(wb[SHEET_SKILLS], SHEET_SKILLS):
    sid = need(row, "编号", SHEET_SKILLS).strip()
    if sid and sid in seen_ids:
        errors.append(f"技能清单 第{row['__row__']}行: 编号 {sid!r} 重复")
    seen_ids.add(sid)

    prompts = split_multi(cell_str(row.get("提示词示例")))
    if not prompts:
        errors.append(f"技能清单 第{row['__row__']}行: 《提示词示例》不能为空")

    status = cell_str(row.get("状态")) or "已上线"
    if status not in ("已上线", "试点", "规划"):
        warnings.append(f"技能清单 第{row['__row__']}行: 状态 {status!r} 不认识，按「已上线」处理")
        status = "已上线"

    skills.append({
        "id": sid,
        "name": need(row, "技能名称", SHEET_SKILLS),
        "en": cell_str(row.get("英文名称")),
        "desc": need(row, "一句话简介", SHEET_SKILLS),
        "biz": need(row, "业务分类", SHEET_SKILLS),
        "func": need(row, "功能分类", SHEET_SKILLS),
        "ownerBiz": need(row, "业务负责人", SHEET_SKILLS),
        "ownerIT": need(row, "IT负责人", SHEET_SKILLS),
        "scenario": cell_str(row.get("适用场景")),
        "prompts": prompts,
        "value": need(row, "提升价值", SHEET_SKILLS),
        "saved": to_int(row.get("单次省时(分钟)"), "单次省时(分钟)", f"技能清单 第{row['__row__']}行"),
        "weekly": to_float(row.get("每周次数"), "每周次数", f"技能清单 第{row['__row__']}行"),
        "tags": split_tags(cell_str(row.get("标签"))),
        "status": status,
        "featured": cell_str(row.get("推荐位")).lower() in ("是", "y", "yes", "true", "1"),
        "weight": to_int(row.get("排序权重"), "排序权重", f"技能清单 第{row['__row__']}行"),
        "row": row["__row__"],
    })

for s in skills:
    s["gain"] = gain_label(s["saved"])
    s["weeklySaved"] = round(s["saved"] * s["weekly"])
    s["monthly"] = round(s["saved"] * s["weekly"] * 4.3)

skills.sort(key=lambda s: (-s["weight"], s["row"]))
for i, s in enumerate(skills, start=1):
    s["index"] = i

# --- 价值案例
CASE_COLS = ["案例编号", "案例标题", "关联技能", "目标角色", "传统做法", "现在一句话",
             "传统耗时(分钟)", "现在耗时(分钟)", "价值说明", "首页展示"]
warn_missing_columns(wb[SHEET_CASES], SHEET_CASES, CASE_COLS)

cases = []
for row in read_rows(wb[SHEET_CASES], SHEET_CASES):
    where = f"价值案例 第{row['__row__']}行"
    before = to_int(row.get("传统耗时(分钟)"), "传统耗时(分钟)", where)
    now = to_int(row.get("现在耗时(分钟)"), "现在耗时(分钟)", where)
    if before <= 0:
        warnings.append(f"{where}: 传统耗时不是正数，耗时对比条会失真")
    steps = [s.strip() for s in re.split(r"[→➜>]+", cell_str(row.get("传统做法"))) if s.strip()]
    cases.append({
        "id": need(row, "案例编号", SHEET_CASES),
        "title": need(row, "案例标题", SHEET_CASES),
        "skill": cell_str(row.get("关联技能")),
        "role": cell_str(row.get("目标角色")),
        "beforeText": cell_str(row.get("传统做法")),
        "beforeSteps": steps,
        "nowPrompt": need(row, "现在一句话", SHEET_CASES),
        "before": before,
        "now": now,
        "value": need(row, "价值说明", SHEET_CASES),
        "home": cell_str(row.get("首页展示")).lower() in ("是", "y", "yes", "true", "1"),
        "row": row["__row__"],
    })
if not cases:
    warnings.append("没有任何价值案例，「价值故事」页会是空的")

# --- 使用指南
guide_steps, guide_faq = [], []
for row in read_rows(wb[SHEET_GUIDE], SHEET_GUIDE):
    kind = cell_str(row.get("类型"))
    item = {"title": cell_str(row.get("标题")), "body": cell_str(row.get("内容"))}
    if kind == "常见问题":
        guide_faq.append(item)
    else:
        guide_steps.append(item)

skills_total_minutes = sum(max(0, c["before"] - c["now"]) for c in cases)

biz_cats: list[str] = []
func_cats: list[str] = []
for s in skills:
    if s["biz"] not in biz_cats:
        biz_cats.append(s["biz"])
    if s["func"] not in func_cats:
        func_cats.append(s["func"])

# 筛选栏里让「用得最多」的排前面
biz_cats.sort(key=lambda c: (-sum(1 for s in skills if s["biz"] == c), c))
func_cats.sort(key=lambda c: (-sum(1 for s in skills if s["func"] == c), c))

# --- 价值页的账：按「单次省时 × 每周次数」汇总
weekly_calls = sum(s["weekly"] for s in skills)
weekly_saved = sum(s["saved"] * s["weekly"] for s in skills)
monthly_saved = round(weekly_saved * 4.3)

domain_map: dict[str, dict] = {}
for s in skills:
    d = domain_map.setdefault(s["biz"], {"name": s["biz"], "count": 0, "weekly": 0.0, "skills": []})
    d["count"] += 1
    d["weekly"] += s["saved"] * s["weekly"]
    d["skills"].append(s["name"])

domains = sorted(domain_map.values(), key=lambda d: -d["weekly"])
for d in domains:
    raw = d["weekly"]
    d["weekly"] = round(raw)
    d["share"] = round(raw / weekly_saved * 100) if weekly_saved else 0

data = {
    "site": {
        "name": cfg["站点名称"],
        "title": cfg["主标题"],
        "desc": cfg["描述"],
        "badge": cfg["角标"],
        "team": cfg["团队署名"],
        "theme": cfg["主题色"],
        "version": cfg["版本"],
        "contact": cfg["联系方式"],
        "footer": cfg["页脚"],
        "about": cfg.get("关于", ""),
    },
    "skills": skills,
    "cases": cases,
    "guide": {"steps": guide_steps, "faq": guide_faq},
    "meta": {
        "generated": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "skillCount": len(skills),
        "caseCount": len(cases),
        "bizCount": len(biz_cats),
        "funcCount": len(func_cats),
        "savedMinutes": skills_total_minutes,
        "bizCategories": biz_cats,
        "funcCategories": func_cats,
        "weeklyCalls": weekly_calls,
        "weeklySavedMin": round(weekly_saved),
        "monthlySavedMin": monthly_saved,
        "manDays": round(monthly_saved / 60 / 8, 1),
        "domains": domains,
    },
}

# ---------------------------------------------------------------- write
if errors:
    print("[ERROR] 数据有问题，已停止生成：")
    for e in errors:
        print("  - " + e)
    if warnings:
        print("\n[WARN]")
        for w in warnings:
            print("  - " + w)
    sys.exit(1)

if DIST.exists():
    shutil.rmtree(DIST)
(DIST / "assets").mkdir(parents=True)

template = (SRC / "template.html").read_text(encoding="utf-8")
replacements = {
    "{{SITE_NAME}}": cfg["站点名称"],
    "{{TITLE}}": cfg["主标题"],
    "{{DESC}}": cfg["描述"][:120],
    "{{TEAM}}": cfg["团队署名"],
    "{{VERSION}}": cfg["版本"],
    "{{FOOTER}}": cfg["页脚"],
    "{{CONTACT}}": cfg["联系方式"],
    "{{ACCENT}}": cfg["主题色"],
    "{{BUILT}}": data["meta"]["generated"],
    "{{COUNT}}": str(len(skills)),
}
for token, value in replacements.items():
    template = template.replace(token, value)

(DIST / "index.html").write_text(template, encoding="utf-8")
shutil.copy2(SRC / "style.css", DIST / "assets" / "style.css")
shutil.copy2(SRC / "app.js", DIST / "assets" / "app.js")

js = "window.SITE_DATA = " + json.dumps(data, ensure_ascii=False, indent=2) + ";\n"
(DIST / "assets" / "data.js").write_text(js, encoding="utf-8")

# ---------------------------------------------------------------- UI 方案预览
THEMES = [
    ("t01", "01 产品蓝（当前）"),
    ("t02", "02 编辑部 · 宋体纸感"),
    ("t03", "03 瑞士 · 黑线极简"),
    ("t04", "04 深空 · 暗色霓虹"),
    ("t05", "05 终端 · 等宽荧光"),
    ("t06", "06 奶咖 · 柔和暖白"),
    ("t07", "07 玻璃 · 渐变磨砂"),
    ("t08", "08 看板 · 紧凑数据"),
    ("t09", "09 宣纸 · 国风朱红"),
    ("t10", "10 粗野 · 黑白硬投影"),
]

PREVIEW_BAR = """
<div id="skinbar">
  <span>UI 方案</span>
  <select id="skinpick" aria-label="切换 UI 方案">
__OPTIONS__
  </select>
  <a href="index.html">看原版</a>
</div>
<style>
#skinbar{position:fixed;left:18px;bottom:18px;z-index:300;display:flex;align-items:center;gap:8px;padding:8px 10px;border-radius:12px;background:rgba(12,17,26,.93);color:#cfd8e3;font:600 12px/1 "Microsoft YaHei",sans-serif;box-shadow:0 18px 40px -18px rgba(0,0,0,.7)}
#skinbar span{color:#8fa0b5;letter-spacing:.06em}
#skinbar select{appearance:none;background:#1b2432;color:#eaf0f8;border:1px solid #2c3a4d;border-radius:8px;padding:7px 10px;font:inherit;cursor:pointer;outline:none}
#skinbar select:hover{border-color:#3f5470}
#skinbar a{color:#9fb0c9;text-decoration:none;padding:6px 4px}
#skinbar a:hover{color:#fff}
@media print{#skinbar{display:none}}
</style>
<script>
(function(){
  var link=document.getElementById('skin'), sel=document.getElementById('skinpick');
  function apply(id){
    link.setAttribute('href','themes/'+id+'.css');
    sel.value=id;
    try{ history.replaceState(null,'','?skin='+id+location.hash); }catch(e){}
  }
  sel.addEventListener('change',function(){ apply(sel.value); });
  var m=location.search.match(/skin=(t\\d\\d)/);
  if(m) apply(m[1]);
})();
</script>
"""

theme_dir = SRC / "themes"
if theme_dir.exists():
    (DIST / "themes").mkdir(exist_ok=True)
    for f in sorted(theme_dir.glob("t*.css")):
        shutil.copy2(f, DIST / "themes" / f.name)

    options = "\n".join(
        f'    <option value="{tid}">{label}</option>' for tid, label in THEMES
    )
    preview = template.replace(
        "</head>",
        '<link id="skin" rel="stylesheet" href="themes/t01.css">\n</head>',
    )
    preview = preview.replace("</body>", PREVIEW_BAR.replace("__OPTIONS__", options) + "\n</body>")
    (DIST / "ui-preview.html").write_text(preview, encoding="utf-8")

total = sum(f.stat().st_size for f in DIST.rglob("*") if f.is_file())
print("[OK] 生成完成")
print(f"     技能 {len(skills)} 个 / 业务分类 {len(biz_cats)} 个 / 功能分类 {len(func_cats)} 个 / 案例 {len(cases)} 个")
print(f"     单次合计省时 {skills_total_minutes} 分钟（约 {skills_total_minutes / 60:.1f} 小时）")
print(f"     每周 {weekly_calls:g} 次调用 / 每周省 {weekly_saved / 60:.1f} 小时 / 每月省 {monthly_saved / 60:.1f} 小时（约 {monthly_saved / 60 / 8:.1f} 人天）")
print(f"     输出 dist/  {total / 1024:.0f} KB")
print(f"     入口 {DIST / 'index.html'}")
for w in warnings:
    print("     [WARN] " + w)
