# -*- coding: utf-8 -*-
"""生成示例数据工作簿《技能展厅.xlsx》。

初次搭建时运行一次即可；日常维护直接编辑 Excel，不需要这个脚本。
想恢复成初始示例：python tools/make_template.py --force
"""

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "技能展厅.xlsx"

HEAD_BG = "1E293B"
HEAD_FG = "FFFFFF"
LINE = Side(style="thin", color="DCE2EA")
BORDER = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)
ZEBRA = PatternFill("solid", fgColor="F7F9FC")

CONFIG = [
    ("站点名称", "Skill 展厅"),
    ("主标题", "「一句话」把专业活儿做完"),
    ("描述", "13 个经过业务验证的智能体技能，覆盖公文、会议、经营、数据与办公交付。找得到、抄得走、用得上——把「这事得干两天」变成「这句话我发你」。"),
    ("角标", "内部智能体平台 · Skill 资产库"),
    ("团队署名", "智能体平台 · 能力建设组"),
    ("主题色", "#2563eb"),
    ("版本", "v1.0"),
    ("联系方式", "skill-platform@example.com"),
    ("页脚", "数据维护在《技能展厅.xlsx》，双击「一键生成.cmd」即可更新本页。"),
]

SKILLS = [
    dict(
        sid="gov-doc-writing", name="公文写作", biz="公文行政", func="写作与提炼",
        owner_biz="张琳", owner_it="王磊",
        desc="按单位、文种、角色和现有材料，生成通知、报告、请示等规范公文，打开就能改。",
        scenario="接到「发个通知」「写份请示」的临时任务，手上只有零散材料；需要套用单位固定版式。",
        prompts=[
            "帮我写一份《关于开展2026年度安全生产大检查的通知》，单位是XX集团，发文对象为各子公司，重点检查消防、用电、危化品三项，10月30日前完成自查并上报。",
            "根据下面这段会议记录，整理成向上级请示的公文，请示事项是追加专项预算80万元，语言正式，结尾用「妥否，请批示」。",
        ],
        value="从「找模板—改抬头—调格式—核口径」的半日工，变成一次对话拿到可编辑成稿。", tags="公文,通知,请示,红头", status="已上线", top=True, weight=100,
    ),
    dict(
        sid="meeting-recap", name="会议纪要", biz="会议协作", func="写作与提炼",
        owner_biz="陈晓", owner_it="王磊",
        desc="把零散发言、录音转写整理成结构化纪要：议题、结论、行动项、负责人、截止日期。",
        scenario="周会、评审会长纪要；会上说法乱、结论散，会后还要挨个追行动项。",
        prompts=[
            "把下面的会议记录整理成纪要，按「议题—讨论要点—决议—行动项（负责人/截止时间）」输出，没提到的信息不要编。",
            "这是1小时的评审会转写稿，帮我提炼5条关键决议和3个待办，每条标注责任人。",
        ],
        value="纪要从「回听录音两小时」变成「贴进去一分钟」，行动项再也不会漏在人名下。", tags="会议,纪要,行动项", status="已上线", top=True, weight=95,
    ),
    dict(
        sid="work-report-writer", name="工作汇报", biz="经营管理", func="写作与提炼",
        owner_biz="陈晓", owner_it="李静",
        desc="按手头记录整理日报、周报、月报，自动核对交付、数字、风险与下期计划。",
        scenario="周五下班前拼周报；月度述职要同时讲清数据、风险和下一步。",
        prompts=[
            "根据我这周的5条工作记录，写一份周报，分「本周交付/关键数字/风险与求助/下周计划」四块，控制在400字内，别堆形容词。",
        ],
        value="周报从「翻聊天记录凑字数」变成「喂记录出结构」，口径统一、领导好读。", tags="周报,月报,述职", status="已上线", top=False, weight=80,
    ),
    dict(
        sid="research-writer", name="研究报告", biz="市场研究", func="写作与提炼",
        owner_biz="李静", owner_it="周涛",
        desc="从研究问题到证据、来源、判断与风险，产出每句结论都带出处的研究报告。",
        scenario="新市场、新赛道摸底，需要一份能直接进决策会讨论的材料。",
        prompts=[
            "围绕「县域咖啡市场是否值得进入」写一份研究报告：市场规模、竞争格局、进入方式三选一、主要风险，每个结论后面标注数据来源。",
        ],
        value="把「搜三天资料还写不出结论」变成「两小时出带出处的初稿」。", tags="研究,调研,决策", status="试点", top=False, weight=70,
    ),
    dict(
        sid="data-viz-gen", name="数据图表", biz="数据洞察", func="数据分析",
        owner_biz="赵敏", owner_it="周涛",
        desc="按数据类型和业务问题选图、配色、标注，直接给出能进汇报的可编辑图表。",
        scenario="汇报要放图，纠结柱状还是折线；配色土、标注乱、重点不突出。",
        prompts=[
            "这是各区域季度销售额，帮我选最合适的图型，生成带同比标注的图表，配色以公司蓝为主色，重点标出下滑的两个区。",
        ],
        value="图表从「配色试半天」变成「说清问题就出图」，一版过会。", tags="图表,可视化,汇报", status="已上线", top=True, weight=90,
    ),
    dict(
        sid="saas-metrics-coach", name="经营指标", biz="经营管理", func="数据分析",
        owner_biz="赵敏", owner_it="周涛",
        desc="核对收入、MRR、客户留存、增长假设与计算口径，指出数字背后的真问题。",
        scenario="月度经营会前，三张表口径对不上，同比环比算法各说各话。",
        prompts=[
            "这是我们本月的收入、MRR、客户数与流失明细，帮我核对口径，算清MRR增速、净留存和获客成本，指出三个最该关注的异常。",
        ],
        value="把「三张表对不上、会上被问住」变成「一套口径、一页结论」。", tags="经营分析,MRR,留存", status="已上线", top=True, weight=88,
    ),
    dict(
        sid="dataset-health-audit", name="数据体检", biz="数据洞察", func="数据分析",
        owner_biz="王磊", owner_it="赵敏",
        desc="检查缺失、异常、单位、口径、重复与可复现计算，先体检再分析。",
        scenario="拿到的业务数据不敢直接用，又说不清到底哪里有问题。",
        prompts=[
            "帮我体检这份销售明细：缺失率、异常值、单位是否统一、有无重复订单，输出问题清单和修复建议。",
        ],
        value="避免「分析三天，结论因脏数据作废」这种最贵的返工。", tags="数据质量,清洗,校验", status="已上线", top=False, weight=60,
    ),
    dict(
        sid="office-docx", name="Word 文档", biz="办公交付", func="文档生成",
        owner_biz="陈晓", owner_it="王磊",
        desc="生成、修改、校验 Word：保留样式、页眉页脚、编号与图表，交付即可用。",
        scenario="要交的是可编辑的正式文档，不是一段聊天文字。",
        prompts=[
            "把这份内容做成Word：A4、仿宋三号、一级标题黑体、带页码页脚，输出可编辑文件。",
        ],
        value="交付物直接能发，不用再手工排版半小时。", tags="Word,排版,交付", status="已上线", top=False, weight=75,
    ),
    dict(
        sid="office-xlsx", name="Excel 表格", biz="办公交付", func="文档生成",
        owner_biz="赵敏", owner_it="李静",
        desc="建表、写公式、做透视与图表，保留原生公式，后续还能自己改。",
        scenario="手工拉公式、做透视表、一格一格调格式。",
        prompts=[
            "把这份流水做成月度汇总表，用SUMIFS按部门和月份汇总，加同比增长列和柱状图。",
        ],
        value="表格从「拉公式拉到眼花」变成「说需求出成品」。", tags="Excel,公式,透视表", status="已上线", top=False, weight=72,
    ),
    dict(
        sid="office-pptx", name="PPT 演示", biz="办公交付", func="文档生成",
        owner_biz="陈晓", owner_it="李静",
        desc="按模板生成结构清晰的汇报 PPT，版式统一、每页一个观点、可直接改。",
        scenario="临时被通知明天要汇报，晚上才开始做 PPT。",
        prompts=[
            "用公司模板做一份10页的季度复盘PPT：结论先行、每页一个观点、数据用图表讲，别写大段文字。",
        ],
        value="从「熬夜排版」到「半小时出可讲的稿」。", tags="PPT,汇报,模板", status="已上线", top=True, weight=85,
    ),
    dict(
        sid="office-export-pdf", name="文档转 PDF", biz="办公交付", func="格式转换",
        owner_biz="李静", owner_it="王磊",
        desc="Word / Excel / PPT 转 PDF，版式不乱、字体不丢、页码不断。",
        scenario="对外发送必须 PDF，转完格式全乱、表格跨页断行。",
        prompts=[
            "把这个Word转成PDF，检查页码和表格有没有跨页断裂，有问题先告诉我。",
        ],
        value="对外交付不再「转一次修一次」。", tags="PDF,转换,对外", status="已上线", top=False, weight=50,
    ),
    dict(
        sid="office-render", name="文档预览", biz="办公交付", func="质量校验",
        owner_biz="李静", owner_it="周涛",
        desc="把文档渲染成图片，发出去之前先看效果，版式问题一眼发现。",
        scenario="文件发出去才发现文字溢出、图片错位，只能道歉重发。",
        prompts=[
            "把这个PPT每页渲染成图片，我检查一下有没有文字溢出和错位。",
        ],
        value="交付前自查一分钟，省掉一轮返工和一次尴尬。", tags="预览,渲染,质检", status="试点", top=False, weight=45,
    ),
    dict(
        sid="office-recalculate", name="表格重算", biz="办公交付", func="质量校验",
        owner_biz="王磊", owner_it="赵敏",
        desc="刷新公式缓存并核对计算值，避免「一打开就是错的数」。",
        scenario="表格公式结果和手算对不上，或者打开一片 0。",
        prompts=[
            "重算这个工作簿的所有公式，把结果和手工核算的差异逐项列出来。",
        ],
        value="数字可信，汇报时不被追问第二遍。", tags="公式,重算,校验", status="规划", top=False, weight=40,
    ),
]

CASES = [
    dict(
        cid="C01", title="两小时录音，变成一分钟纪要", skill="会议纪要", role="项目助理",
        before="回听录音3遍→手打讨论要点→群里逐个确认口径→整理成行动项表格",
        now="把1小时评审会转写稿整理成纪要：议题、决议、行动项（负责人/截止时间），没提到的别编。",
        bmin=120, nmin=1,
        value="会后10分钟发出纪要，5个行动项直接落到人头，下一次会不用再对上次的账。",
        home=True,
    ),
    dict(
        cid="C02", title="三张表对不上，变成一页结论", skill="经营指标", role="经营分析",
        before="三个系统各导一张表→手工对齐口径→Excel透视→算同比→会上被问住现场重算",
        now="核对本月收入、MRR、客户与流失明细的口径，算清增速、净留存、获客成本，指出三个异常。",
        bmin=180, nmin=3,
        value="经营会从「解释数字」变成「讨论动作」，口径争议当场消掉。",
        home=True,
    ),
    dict(
        cid="C03", title="半天的通知，变成一句话", skill="公文写作", role="行政 / 文秘",
        before="翻去年同类文→改抬头改日期→调字体行距→找人核口径→打印校对三遍",
        now="写一份《安全生产大检查通知》，单位XX集团，对象各子公司，查消防/用电/危化品，10月30日前上报。",
        bmin=240, nmin=5,
        value="当天下午就能发文，格式一次过，文秘不用再做「排版工」。",
        home=True,
    ),
    dict(
        cid="C04", title="配色试五版，变成说清问题就出图", skill="数据图表", role="业务汇报人",
        before="挑图型→配色试5版→手工加标注→调坐标轴→导出贴进PPT",
        now="用各区域季度销售额选最合适的图，带同比标注，公司蓝为主色，突出下滑的两个区。",
        bmin=90, nmin=2,
        value="汇报图一眼看出重点，领导不再问「所以呢」。",
        home=False,
    ),
    dict(
        cid="C05", title="周五拼周报，变成喂记录出结构", skill="工作汇报", role="一线主管",
        before="翻聊天记录和邮件→拼凑本周条目→补数字→统一措辞→反复改语气",
        now="根据这周5条工作记录写周报，分交付/数字/风险/下周计划，400字内。",
        bmin=90, nmin=2,
        value="周五下午不再耗在措辞上，周报还能看出趋势。",
        home=False,
    ),
]

GUIDE = [
    ("步骤", "找到你要的能力", "在技能库按业务分类、功能分类或关键词筛选；卡片上的「复制提示词」可以直接用。"),
    ("步骤", "把示例改成你的场景", "复制后替换括号里的业务信息（单位、时间、数据口径）。写得越具体，产出越接近成稿。"),
    ("步骤", "交给智能体执行", "粘贴发送即可。需要交付文件时，直接说「输出 Word / Excel / PPT」。"),
    ("步骤", "有问题找谁", "卡片上标了业务负责人（提需求、验收效果）和 IT 负责人（维护提示词、修问题）。"),
    ("常见问题", "网页需要联网吗？", "不需要。整个站点是本地文件，双击 index.html 就能看，断网、内网、发给别人都一样。"),
    ("常见问题", "提示词能直接改吗？", "能。示例是起点不是标准答案，把占位信息换成真实业务，效果会明显更好。"),
    ("常见问题", "数据在哪里维护？", "全部在《技能展厅.xlsx》。改完双击「一键生成.cmd」，dist 目录里的网页就更新了。"),
    ("常见问题", "业务分类和功能分类有什么区别？", "业务分类是「哪个业务域在用」，功能分类是「这个能力干什么活」。两者交叉，能快速定位。"),
    ("常见问题", "能发给外部或客户看吗？", "可以，但请先清掉内部联系方式、未公开的业务数据和试点中的技能。"),
]

HELP = [
    ("站点配置", "配置项 / 值", "是", "键值对。站点名称、主标题、描述、角标、团队署名、主题色、版本、联系方式、页脚。主标题里用「」包住的字会高亮显示。"),
    ("技能清单", "编号", "是", "英文短横线命名，唯一，作为详情页链接和去重依据。"),
    ("技能清单", "技能名称", "是", "卡片标题，4-8 个字最好看。"),
    ("技能清单", "英文名称", "否", "卡片的英文副标题、弹窗里的英文名，如 Meeting Recap；留空则不显示。"),
    ("技能清单", "一句话简介", "是", "说清「它能替你做什么」，控制在 40 字内，卡片会截断两行。"),
    ("技能清单", "业务分类", "是", "从业务视角分域，如 公文行政 / 会议协作 / 经营管理 / 市场研究 / 数据洞察 / 办公交付。建议 5-6 个。"),
    ("技能清单", "功能分类", "是", "这个能力干的活，如 写作与提炼 / 数据分析 / 文档生成 / 格式转换 / 质量校验。建议 4-5 个。"),
    ("技能清单", "业务负责人", "是", "提需求、验收效果的人；卡片和详情里显示。"),
    ("技能清单", "IT负责人", "是", "维护提示词、修问题、做迭代的人；卡片和详情里显示。"),
    ("技能清单", "适用场景", "否", "什么时候该用它，写在详情抽屉里。"),
    ("技能清单", "提示词示例", "是", "可直接复制的完整指令。多个示例用单独一行 --- 分隔，页面会拆成多个复制块。"),
    ("技能清单", "提升价值", "是", "一句话讲清价值，最好带数字或对比。"),
    ("技能清单", "单次省时(分钟)", "否", "纯数字。一次任务能省多少分钟，用于算账和生成「单次省 X」标签；留空则该技能不进汇总。"),
    ("技能清单", "每周次数", "否", "纯数字，可以是小数（0.25 = 每月一次）。写业务真实节奏，价值页用它算每周/每月节省。"),
    ("技能清单", "标签", "否", "逗号分隔，参与搜索。"),
    ("技能清单", "状态", "否", "已上线 / 试点 / 规划，三色标识，默认已上线。"),
    ("技能清单", "推荐位", "否", "预留字段：填「是」会在首页优先展示。"),
    ("技能清单", "排序权重", "否", "数字越大越靠前，同分按行序。"),
    ("价值案例", "案例编号", "是", "C01、C02… 递增即可。"),
    ("价值案例", "案例标题", "是", "用「以前 X，现在 Y」的句式最有冲击力。"),
    ("价值案例", "关联技能", "是", "填技能名称或编号。"),
    ("价值案例", "目标角色", "否", "这个案例站在谁的视角，如「项目助理」。"),
    ("价值案例", "传统做法", "是", "用 → 分隔步骤，页面会渲染成步骤清单并逐条播放；数字要真实。"),
    ("价值案例", "现在一句话", "是", "实际用的提示词，会以打字效果呈现，并原样提供给读者复制。"),
    ("价值案例", "传统耗时(分钟)", "是", "纯数字，用于计数动画和对比条。"),
    ("价值案例", "现在耗时(分钟)", "是", "纯数字，通常 1-5 分钟。"),
    ("价值案例", "价值说明", "是", "省下来的时间换回了什么，不要只写「效率提升」。"),
    ("价值案例", "首页展示", "否", "预留字段：首页现在只放技能库，这一列暂不生效。"),
    ("使用指南", "类型 / 标题 / 内容", "是", "类型填「步骤」或「常见问题」，决定它出现在指南页哪一区。"),
]


EN = {
    "gov-doc-writing": "Official Document Writer",
    "meeting-recap": "Meeting Recap",
    "work-report-writer": "Work Report",
    "research-writer": "Research Report",
    "data-viz-gen": "Chart Builder",
    "saas-metrics-coach": "Business Metrics",
    "dataset-health-audit": "Data Health Check",
    "office-docx": "Word Document",
    "office-xlsx": "Excel Workbook",
    "office-pptx": "Slide Deck",
    "office-export-pdf": "PDF Export",
    "office-render": "Doc Preview",
    "office-recalculate": "Sheet Recalculate",
}


# 单次节省分钟数 / 每周使用次数（用于价值页算账）
SAVED = {
    "gov-doc-writing": 235,
    "meeting-recap": 119,
    "work-report-writer": 88,
    "research-writer": 360,
    "data-viz-gen": 88,
    "saas-metrics-coach": 177,
    "dataset-health-audit": 120,
    "office-docx": 40,
    "office-xlsx": 60,
    "office-pptx": 240,
    "office-export-pdf": 20,
    "office-render": 15,
    "office-recalculate": 30,
}

WEEKLY = {
    "gov-doc-writing": 2,
    "meeting-recap": 3,
    "work-report-writer": 1,
    "research-writer": 0.25,
    "data-viz-gen": 2,
    "saas-metrics-coach": 0.25,
    "dataset-health-audit": 0.5,
    "office-docx": 1,
    "office-xlsx": 1,
    "office-pptx": 0.5,
    "office-export-pdf": 2,
    "office-render": 1,
    "office-recalculate": 0.5,
}


def style_sheet(ws, headers, rows, widths):
    ws.freeze_panes = "A2"
    for i, name in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = Font(name="微软雅黑", size=10.5, bold=True, color=HEAD_FG)
        c.fill = PatternFill("solid", fgColor=HEAD_BG)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.row_dimensions[1].height = 30
    for r, row in enumerate(rows, start=2):
        for cidx, value in enumerate(row, start=1):
            c = ws.cell(row=r, column=cidx, value=value)
            c.font = Font(name="微软雅黑", size=10)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = BORDER
            if r % 2 == 1:
                c.fill = ZEBRA
        ws.row_dimensions[r].height = 42
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"


def main():
    if OUT.exists() and "--force" not in sys.argv:
        print(f"[SKIP] {OUT.name} 已存在。要覆盖（会丢失现有数据）请加 --force")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "站点配置"
    style_sheet(ws, ["配置项", "值"], CONFIG, [18, 100])

    ws = wb.create_sheet("技能清单")
    rows = [[
        s["sid"], s["name"], EN[s["sid"]], s["desc"], s["biz"], s["func"], s["owner_biz"], s["owner_it"],
        s["scenario"], "\n---\n".join(s["prompts"]), SAVED[s["sid"]], WEEKLY[s["sid"]],
        s["value"], s["tags"], s["status"], "是" if s["top"] else "否", s["weight"],
    ] for s in SKILLS]
    style_sheet(
        ws,
        ["编号", "技能名称", "英文名称", "一句话简介", "业务分类", "功能分类", "业务负责人", "IT负责人",
         "适用场景", "提示词示例", "单次省时(分钟)", "每周次数", "提升价值", "标签", "状态", "推荐位", "排序权重"],
        rows,
        [18, 15, 26, 44, 12, 12, 11, 10, 38, 60, 14, 10, 38, 20, 10, 9, 10],
    )
    dv_status = DataValidation(type="list", formula1='"已上线,试点,规划"', allow_blank=True)
    dv_status.error = "只能填 已上线 / 试点 / 规划"
    ws.add_data_validation(dv_status)
    dv_status.add(f"O2:O{len(rows) + 1}")
    dv_top = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
    ws.add_data_validation(dv_top)
    dv_top.add(f"P2:P{len(rows) + 1}")

    ws = wb.create_sheet("价值案例")
    rows = [[
        c["cid"], c["title"], c["skill"], c["role"], c["before"], c["now"],
        c["bmin"], c["nmin"], c["value"], "是" if c["home"] else "否",
    ] for c in CASES]
    style_sheet(
        ws,
        ["案例编号", "案例标题", "关联技能", "目标角色", "传统做法", "现在一句话",
         "传统耗时(分钟)", "现在耗时(分钟)", "价值说明", "首页展示"],
        rows,
        [10, 34, 14, 14, 52, 52, 14, 14, 40, 10],
    )
    dv_home = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
    ws.add_data_validation(dv_home)
    dv_home.add(f"J2:J{len(rows) + 1}")

    ws = wb.create_sheet("使用指南")
    rows = [[i, kind, title, body] for i, (kind, title, body) in enumerate(GUIDE, start=1)]
    style_sheet(ws, ["序号", "类型", "标题", "内容"], rows, [8, 12, 26, 88])

    ws = wb.create_sheet("填写说明")
    style_sheet(ws, ["Sheet", "字段", "是否必填", "填写说明"], HELP, [12, 20, 10, 100])

    wb.save(OUT)
    print(f"[OK] 已生成 {OUT}")
    print(f"     技能 {len(SKILLS)} 个 / 案例 {len(CASES)} 个")


if __name__ == "__main__":
    main()
